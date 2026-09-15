#!/usr/bin/env python3
"""One-off bootstrap: seed the item master and the 2025 prices from the office
manager's "Acme Widget Top 25 Items Comparison 2025" working workbook.

Implements spec section 5 (docs/spec.md). Stdlib + openpyxl only - deliberately
does NOT import the `supplytrack` package, so it can run before the package
exists and never drifts with it. The key-building functions below are duplicated
from the spec on purpose and must stay byte-identical to `supplytrack.ingest`.

Writes exactly two files:
    <data-dir>/item_master.csv
    <data-dir>/2025/prices.csv
and a verification summary to stdout (optionally also to --summary).
Nothing else on disk is touched; the workbook is opened read-only.
"""

from __future__ import annotations

import argparse
import collections
import csv
import re
import sys
import unicodedata
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles.fonts import Font

# The 2025 workbooks carry <font family="34">, outside openpyxl's permitted
# range (max 14). Relax the descriptor bound before any load_workbook call.
Font.family.max = 99

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent.parent  # src/scripts -> src -> project root

DEFAULT_INBOX = PROJECT / "inbox" / "Acme Widget Top 25 Items Comparison 2025.xlsx"
DEFAULT_DATA = PROJECT / "data"

YEAR = "2025"
CHECKED_ON = "2026-03-26"
VENDORS = ["Office Depot", "Preferred", "Amazon", "Staples"]

SHEET_RAW = "orders_from_20250101_to_2025123"
SHEET_ONLY = "AMZ Office Supply Orders ONLY"
SHEET_PBS = "PBS Orders"
SHEET_AGG = "Sheet3"
SHEET_REPORT = "Sheet1"

MASTER_COLUMNS = [
    "key", "source", "raw_title", "include", "canonical_name", "units_per_pack",
    "unit_label", "upp_source", "amazon_category", "first_seen", "last_seen", "note",
]
PRICE_COLUMNS = [
    "rank", "canonical_name", "vendor", "unit_price", "status", "url", "checked_on", "note",
]

# --------------------------------------------------------------------------
# Item identity (spec section 3). Keep in step with supplytrack.ingest.
# --------------------------------------------------------------------------


def normalise_title(title: str) -> str:
    t = unicodedata.normalize("NFKC", title or "").casefold()
    t = re.sub(r"[^0-9a-z /\-.()]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def amazon_key(title: str) -> str:
    return "amz:" + normalise_title(title)


def preferred_key(code: str) -> str:
    return "pbs:" + (code or "").strip().upper()


# --------------------------------------------------------------------------
# Units per pack
# --------------------------------------------------------------------------

_PACK_RE = re.compile(r"^\s*([A-Za-z]+)\s*(\d+)\s*$")


def decode_preferred_pack(pack: str) -> int | None:
    """CT10 -> 10, BX100 -> 100, CT2500 -> 2500, CS1 -> 1, Each -> 1."""
    p = (pack or "").strip()
    if not p:
        return None
    if p.casefold() in {"each", "ea"}:
        return 1
    m = _PACK_RE.match(p)
    if m:
        return int(m.group(2))
    return None


# Ordered so the more specific phrase wins the `phrase` slot for a given value.
_UPP_PATTERNS = [
    r"\(\s*bulk\s+(\d[\d,]*)\s*[- ]?pack\s*\)",   # (Bulk 1000 Pack)
    r"\bpack\s+of\s+(\d[\d,]*)\b",                 # Pack of 12
    r"\bbox\s+of\s+(\d[\d,]*)\b",                  # Box of 100
    r"\b(\d[\d,]*)\s*/\s*pack\b",                  # 48/Pack
    r"\b(\d[\d,]*)\s*[- ]?pack\b",                 # 6 Pack / 24-Pack
    r"\b(\d[\d,]*)\s*[- ]?count\b",                # 32 Count / 100-Count
    r"\b(\d[\d,]*)\s*[- ]?ct\b",                   # 40 Ct
    r"\b(\d[\d,]*)\s*[- ]?pk\b",                   # 12 pk
    r"\b(\d[\d,]*)\s*pads\b",                      # 6 Pads
    r"\b(\d[\d,]*)\s*sheets\b",                    # 500 Sheets
]
_UPP_RES = [re.compile(p, re.IGNORECASE) for p in _UPP_PATTERNS]


def upp_candidates(title: str) -> list[tuple[int, str]]:
    """Distinct (value, phrase) pack-size candidates read out of a product title."""
    found: dict[int, str] = {}
    for rx in _UPP_RES:
        for m in rx.finditer(title or ""):
            try:
                value = int(m.group(1).replace(",", ""))
            except (TypeError, ValueError):
                continue
            if value <= 0:
                continue
            found.setdefault(value, " ".join(m.group(0).split()))
    return sorted(found.items())


def format_candidates(cands: list[tuple[int, str]]) -> str:
    if not cands:
        return ""
    return "candidates: " + "|".join("%d (%s)" % (v, p) for v, p in cands)


# --------------------------------------------------------------------------
# The 2025 workbook's hand-made merges and corrections (analysis A9.6, spec s5)
# --------------------------------------------------------------------------

# Sheet3 row -> extra Amazon titles folded into it by hand.
VARIANT_MERGES: dict[int, list[str]] = {
    4: ["Energizer AA Batteries, Alkaline Power Double A Battery Alkaline, 32 Count"],
    24: [
        "Germ-X Original Hand Sanitizer, Kids Hand Sanitizer, Non-Drying Moisturizing "
        "Gel with Vitamin E, Instant and No Rinse Formula, Bulk Mini Travel Size for "
        "On-The-Go, 2 Fl Oz (Display Pack of 6)"
    ],
    29: [
        'Five Star Spiral Notebook + Study App, 6 Pack, 1 Subject, College Ruled Paper, '
        '8-1/2" x 11", 100 Sheets, Fights Ink Bleed, Water Resistant Cover, Assorted '
        'Colors (38052)'
    ],
    40: [
        'Cambridge Limited Business Notebook, Legal Ruled, 8-1/4" x 11", 80 Sheets, '
        'Soft Touch Flexible Cover, Wirebound, Gray (06062)'
    ],
}

# Sheet3 row -> the row whose label becomes the shared canonical name.
CANONICAL_ALIAS: dict[int, int] = {15: 9}  # both birthday-card listings -> Harloon

# Preferred codes whose Description does not equal any Sheet3 label verbatim.
PREFERRED_ROW_OVERRIDE: dict[str, int] = {"BSN36591": 7, "BSN36590": 20}

EXCLUDE_NOTE = "excluded by the office manager in 2025 (not office supply)"

# Sheet3 row -> per-key corrections applied after the generic derivation.
ROW_OVERRIDES: dict[int, dict] = {
    3: {"include": "n", "units_per_pack": "", "upp_source": "", "note": EXCLUDE_NOTE},
    23: {"include": "n", "units_per_pack": "", "upp_source": "", "note": EXCLUDE_NOTE},
    6: {
        "units_per_pack": 36, "upp_source": "title",
        "note": "2025 workbook count 288 does not reconcile to export (10 packs x 28.8); "
                "title says 36 Boxes",
    },
    14: {
        "units_per_pack": 400, "upp_source": "title",
        "note": "2025 workbook used 100; title says 400",
    },
    24: {
        "units_per_pack": 6, "upp_source": "title",
        "note": "2025 workbook used 12; title says Display Pack of 6",
    },
    9: {
        "units_per_pack": 200, "upp_source": "2025-workbook",
        "note": "merged in 2025 workbook; 2025 report showed 200, actual 300",
    },
    15: {
        "units_per_pack": 100, "upp_source": "2025-workbook",
        "note": "merged in 2025 workbook; 2025 report showed 200, actual 300",
    },
}

# Rows the reconciliation is expected to disagree on (documented in the analysis).
DOCUMENTED_RECON_ROWS = {
    3: "zip ties - excluded by hand",
    4: "Energizer AA - A9.6, B left at 8 but C reflects the merged variant",
    6: "Kleenex - A9.6, B=10 vs pivot 11, C/B non-integer",
    14: "Avery 5168 - A9.2, title says 400",
    23: "ARTISTRO - excluded by hand",
    24: "Germ-X - A9.2, title says 6",
}


# --------------------------------------------------------------------------
# Workbook reading
# --------------------------------------------------------------------------


def header_index(ws) -> dict[str, int]:
    return {
        str(c.value).strip().casefold(): c.column
        for c in ws[1]
        if c.value is not None
    }


def read_amazon(ws) -> list[dict]:
    idx = header_index(ws)
    need = ["title", "item quantity", "amazon-internal product category"]
    missing = [n for n in need if n not in idx]
    if missing:
        raise SystemExit("%s: missing columns %s" % (ws.title, missing))
    rows = []
    for r in range(2, ws.max_row + 1):
        title = ws.cell(r, idx["title"]).value
        if title is None or str(title).strip() == "":
            continue
        qty = ws.cell(r, idx["item quantity"]).value or 0
        cat = ws.cell(r, idx["amazon-internal product category"]).value or ""
        rows.append({"title": str(title), "qty": qty, "category": str(cat).strip()})
    return rows


def read_preferred(ws) -> list[dict]:
    idx = header_index(ws)
    rows = []
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, idx["code"]).value
        if code is None or str(code).strip() == "":
            continue
        rows.append({
            "code": str(code).strip(),
            "description": str(ws.cell(r, idx["description"]).value or "").strip(),
            "pack": str(ws.cell(r, idx["pack"]).value or "").strip(),
            "quan": ws.cell(r, idx["quan"]).value or 0,
        })
    return rows


def read_sheet3(ws) -> list[dict]:
    rows = []
    for r in range(2, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        if str(label).strip().casefold() == "grand total":
            continue
        rows.append({
            "row": r,
            "label": str(label),
            "b": ws.cell(r, 2).value,
            "c": ws.cell(r, 3).value,
            "d": str(ws.cell(r, 4).value or "").strip(),
        })
    return rows


def number_as_typed(value) -> str:
    """Render a numeric cell without float artefacts."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        d = Decimal(str(value)).normalize()
        if d == d.to_integral_value():
            d = d.quantize(Decimal(1))
        return format(d, "f")
    return str(value)


# --------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------


def build(inbox: Path, data_dir: Path) -> str:
    out: list[str] = []

    def say(line: str = "") -> None:
        out.append(line)

    wb = openpyxl.load_workbook(inbox, data_only=True)
    for name in (SHEET_RAW, SHEET_ONLY, SHEET_PBS, SHEET_AGG, SHEET_REPORT):
        if name not in wb.sheetnames:
            raise SystemExit("%s: sheet %r not found" % (inbox.name, name))

    raw_rows = read_amazon(wb[SHEET_RAW])
    only_rows = read_amazon(wb[SHEET_ONLY])
    pbs_rows = read_preferred(wb[SHEET_PBS])
    agg_rows = read_sheet3(wb[SHEET_AGG])
    report = wb[SHEET_REPORT]

    # ---- Amazon keys from the raw export -------------------------------
    raw_qty: collections.Counter = collections.Counter()
    first_title: dict[str, str] = {}
    categories: dict[str, list[str]] = collections.defaultdict(list)
    exact_titles: dict[str, set[str]] = collections.defaultdict(set)
    for row in raw_rows:
        k = amazon_key(row["title"])
        first_title.setdefault(k, row["title"])
        raw_qty[k] += row["qty"]
        exact_titles[k].add(row["title"])
        if row["category"] and row["category"] not in categories[k]:
            categories[k].append(row["category"])

    only_exact = {row["title"] for row in only_rows}
    only_qty: collections.Counter = collections.Counter()
    for row in only_rows:
        only_qty[amazon_key(row["title"])] += row["qty"]

    # A key must be wholly in or wholly out of the ONLY sheet, else `include`
    # would be arbitrary for the exact titles that collapse onto it.
    mixed_keys = [
        k for k, titles in exact_titles.items()
        if 0 < len(titles & only_exact) < len(titles)
    ]

    included_keys = {k for k in raw_qty if exact_titles[k] & only_exact}

    # ---- Sheet3 -> keys -------------------------------------------------
    agg_by_row = {r["row"]: r for r in agg_rows}
    agg_by_norm = {normalise_title(r["label"]): r["row"] for r in agg_rows}

    row_to_amz: dict[int, list[str]] = collections.defaultdict(list)
    amz_to_row: dict[str, int] = {}
    for r in agg_rows:
        k = amazon_key(r["label"])
        if k in raw_qty:
            row_to_amz[r["row"]].append(k)
            amz_to_row[k] = r["row"]
    for row_no, titles in VARIANT_MERGES.items():
        for title in titles:
            k = amazon_key(title)
            if k not in raw_qty:
                raise SystemExit("merge variant not found in export: %r" % title)
            row_to_amz[row_no].append(k)
            amz_to_row[k] = row_no

    row_to_pbs: dict[int, list[str]] = collections.defaultdict(list)
    pbs_to_row: dict[str, int] = {}
    pbs_codes: dict[str, dict] = {}
    pbs_qty: collections.Counter = collections.Counter()
    for row in pbs_rows:
        k = preferred_key(row["code"])
        pbs_codes.setdefault(k, row)
        pbs_qty[k] += row["quan"]
    for k, row in pbs_codes.items():
        code = k.split(":", 1)[1]
        row_no = PREFERRED_ROW_OVERRIDE.get(code)
        if row_no is None:
            row_no = agg_by_norm.get(normalise_title(row["description"]))
        if row_no is not None:
            row_to_pbs[row_no].append(k)
            pbs_to_row[k] = row_no

    def canonical_for_row(row_no: int) -> str:
        return agg_by_row[CANONICAL_ALIAS.get(row_no, row_no)]["label"]

    # units_per_pack implied by Sheet3, usable only when the row reconciles.
    row_upp: dict[int, int | None] = {}
    for r in agg_rows:
        row_no, b, c = r["row"], r["b"], r["c"]
        observed = (sum(raw_qty[k] for k in row_to_amz.get(row_no, []))
                    + sum(pbs_qty[k] for k in row_to_pbs.get(row_no, [])))
        ok = bool(b) and observed == b
        upp = None
        if ok and isinstance(c, (int, float)) and b and float(c) % float(b) == 0:
            upp = int(float(c) / float(b))
        row_upp[row_no] = upp

    # ---- Master rows ----------------------------------------------------
    master: dict[str, dict] = {}

    for k in raw_qty:
        include = "y" if k in included_keys else "n"
        row_no = amz_to_row.get(k)
        canonical = first_title[k]
        upp: object = ""
        upp_source = ""
        note = ""
        unit_label = "EA"
        if include == "y" and row_no is not None:
            canonical = canonical_for_row(row_no)
            if agg_by_row[row_no]["d"].upper() == "RM":
                unit_label = "RM"
            if row_upp[row_no] is not None:
                upp = row_upp[row_no]
                upp_source = "2025-workbook"
            else:
                cands = upp_candidates(first_title[k])
                if len(cands) == 1:
                    upp = cands[0][0]
                    upp_source = "title"
                note = ("2025 workbook row does not reconcile to the export "
                        "(Sheet3 B=%s); %s" % (agg_by_row[row_no]["b"],
                                               format_candidates(cands))).strip("; ")
        elif include == "y":
            # The M-Z tail Sheet3 never covered: a single unambiguous regex hit
            # is taken, anything else is left blank for the review queue. The
            # source is still the title either way, so `review` can select on it.
            cands = upp_candidates(first_title[k])
            if len(cands) == 1:
                upp = cands[0][0]
            upp_source = "title"
            note = format_candidates(cands)
        master[k] = {
            "key": k,
            "source": "amazon",
            "raw_title": first_title[k],
            "include": include,
            "canonical_name": canonical,
            "units_per_pack": upp,
            "unit_label": unit_label,
            "upp_source": upp_source,
            "amazon_category": categories[k][0] if categories[k] else "",
            "first_seen": YEAR,
            "last_seen": YEAR,
            "note": note,
        }

    pack_conflicts: list[str] = []
    for k, row in pbs_codes.items():
        row_no = pbs_to_row.get(k)
        decoded = decode_preferred_pack(row["pack"])
        canonical = row["description"]
        upp = decoded if decoded is not None else ""
        upp_source = "2025-workbook" if decoded is not None else ""
        note = ""
        if row_no is not None:
            canonical = canonical_for_row(row_no)
            implied = row_upp[row_no]
            if implied is not None:
                upp = implied
                upp_source = "2025-workbook"
                if decoded is not None and decoded != implied:
                    note = ("Pack code %s decodes to %d; 2025 workbook counted %d "
                            "per purchase unit" % (row["pack"], decoded, implied))
                    pack_conflicts.append("%s: %s=%d vs Sheet3 %d"
                                          % (k, row["pack"], decoded, implied))
        master[k] = {
            "key": k,
            "source": "preferred",
            "raw_title": row["description"],
            "include": "y",
            "canonical_name": canonical,
            "units_per_pack": upp,
            "unit_label": "RM" if "ream" in row["description"].casefold() else "EA",
            "upp_source": upp_source,
            "amazon_category": "",
            "first_seen": YEAR,
            "last_seen": YEAR,
            "note": note,
        }

    # ---- Per-row overrides (spec section 5) -----------------------------
    for row_no, override in ROW_OVERRIDES.items():
        keys = list(row_to_amz.get(row_no, [])) + list(row_to_pbs.get(row_no, []))
        if not keys:
            raise SystemExit("override for Sheet3 row %d matched no key" % row_no)
        for k in keys:
            for field, value in override.items():
                master[k][field] = value
            if override.get("include") == "n":
                master[k]["canonical_name"] = master[k]["raw_title"]

    # ---- Write item_master.csv -----------------------------------------
    data_dir.mkdir(parents=True, exist_ok=True)
    master_path = data_dir / "item_master.csv"
    ordered = sorted(master.values(), key=lambda r: (r["source"], r["key"]))
    with master_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(MASTER_COLUMNS)
        for r in ordered:
            w.writerow([r[c] for c in MASTER_COLUMNS])

    # ---- Write 2025/prices.csv ------------------------------------------
    price_rows: list[list] = []
    price_issues: list[str] = []
    for r in range(7, 32):
        rank = report.cell(r, 2).value
        desc = str(report.cell(r, 3).value or "")
        row_no = agg_by_norm.get(normalise_title(desc))
        if row_no is None:
            # Sheet1 row 13 concatenates both birthday-card titles.
            norm_desc = normalise_title(desc)
            hits = {CANONICAL_ALIAS.get(rn, rn)
                    for lbl, rn in agg_by_norm.items() if lbl and lbl in norm_desc}
            if len(hits) != 1:
                raise SystemExit("Sheet1 row %d: cannot map description to Sheet3" % r)
            row_no = hits.pop()
        canonical = canonical_for_row(row_no)
        for offset, vendor in enumerate(VENDORS):
            cell = report.cell(r, 6 + offset).value
            if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                price, status = number_as_typed(cell), "priced"
            else:
                text = str(cell or "").strip()
                price = ""
                low = text.casefold()
                if low == "not available":
                    status = "not_available"
                elif low == "discontinued":
                    status = "discontinued"
                else:
                    status = ""
                    price_issues.append("Sheet1 %s%d: unexpected %r"
                                        % (chr(70 + offset), r, text))
            price_rows.append([rank, canonical, vendor, price, status, "", CHECKED_ON, ""])

    year_dir = data_dir / YEAR
    year_dir.mkdir(parents=True, exist_ok=True)
    prices_path = year_dir / "prices.csv"
    with prices_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(PRICE_COLUMNS)
        w.writerows(price_rows)

    # ---- Verification summary -------------------------------------------
    say("Bootstrap 2025 - verification summary")
    say("workbook : %s" % inbox)
    say("master   : %s" % master_path)
    say("prices   : %s" % prices_path)
    say()
    say("Amazon export lines read      : %d" % len(raw_rows))
    say("AMZ ONLY sheet lines read     : %d" % len(only_rows))
    say("Preferred lines read          : %d" % len(pbs_rows))
    say("Sheet3 item rows              : %d" % len(agg_rows))
    say()
    say("Master rows written           : %d" % len(ordered))
    for src, n in sorted(collections.Counter(r["source"] for r in ordered).items()):
        say("  source=%-10s %d" % (src, n))
    for (src, inc), n in sorted(
            collections.Counter((r["source"], r["include"]) for r in ordered).items()):
        say("  source=%-10s include=%s  %d" % (src, inc, n))
    for src, n in sorted(
            collections.Counter(r["upp_source"] or "(blank)" for r in ordered).items()):
        say("  upp_source=%-16s %d" % (src, n))
    say()

    blank_upp = [r for r in ordered if r["include"] == "y" and r["units_per_pack"] == ""]
    say("Included rows with blank units_per_pack: %d" % len(blank_upp))
    for r in blank_upp:
        say("  %s" % r["key"][:120])
    say()

    by_canonical: dict[str, list[str]] = collections.defaultdict(list)
    for r in ordered:
        by_canonical[r["canonical_name"]].append(r["key"])
    multi = {name: keys for name, keys in by_canonical.items() if len(keys) > 1}
    say("Canonical names carrying more than one key: %d" % len(multi))
    for name, keys in sorted(multi.items()):
        say("  %s" % name[:110])
        for k in keys:
            say("      %s" % k[:120])
    say()

    say("Prices rows written: %d" % len(price_rows))
    for status, n in sorted(
            collections.Counter(r[4] or "(blank)" for r in price_rows).items()):
        say("  status=%-16s %d" % (status, n))
    for issue in price_issues:
        say("  ISSUE %s" % issue)
    missing_canon = sorted({r[1] for r in price_rows} - set(by_canonical))
    say("Prices canonical names absent from the master: %d" % len(missing_canon))
    for name in missing_canon:
        say("  %s" % name[:120])
    say()

    say("Sheet3 reconciliation (raw-export Item Quantity + Preferred Quan vs Sheet3 B)")
    mismatches = []
    for r in agg_rows:
        row_no = r["row"]
        amz_keys = row_to_amz.get(row_no, [])
        pbs_keys = row_to_pbs.get(row_no, [])
        observed = (sum(raw_qty[k] for k in amz_keys)
                    + sum(pbs_qty[k] for k in pbs_keys))
        only_sum = sum(only_qty[k] for k in amz_keys)
        if not amz_keys and not pbs_keys:
            mismatches.append((row_no, r["b"], None, only_sum, "no key matched"))
        elif observed != r["b"]:
            mismatches.append((row_no, r["b"], observed, only_sum, ""))
    say("  rows checked: %d   mismatches: %d" % (len(agg_rows), len(mismatches)))
    for row_no, b, observed, only_sum, why in mismatches:
        tag = DOCUMENTED_RECON_ROWS.get(row_no, "UNDOCUMENTED")
        say("  Sheet3 R%-3d B=%s observed=%s (ONLY sheet=%s) %s [%s]"
            % (row_no, b, observed, only_sum, why, tag))
        say("      %s" % agg_by_row[row_no]["label"][:110])
    say()

    say("Other checks")
    say("  keys mixed across the ONLY sheet (must be 0): %d" % len(mixed_keys))
    for k in mixed_keys:
        say("      %s" % k[:120])
    multi_cat = [k for k, cats in categories.items() if len(cats) > 1]
    say("  keys with more than one Amazon category: %d" % len(multi_cat))
    for k in multi_cat[:10]:
        say("      %s -> %s" % (k[:80], categories[k]))
    say("  Preferred codes with Pack-code vs Sheet3 conflict: %d" % len(pack_conflicts))
    for line in pack_conflicts:
        say("      %s" % line)
    unmapped_pbs = [k for k in pbs_codes if k not in pbs_to_row]
    say("  Preferred codes not mapped to a Sheet3 row: %d" % len(unmapped_pbs))
    for k in unmapped_pbs:
        say("      %s" % k)
    tail = sorted(included_keys - set(amz_to_row))
    say("  Included Amazon titles absent from Sheet3 (the M-Z tail): %d" % len(tail))
    say()

    return "\n".join(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--inbox", type=Path, default=DEFAULT_INBOX,
                    help="the 2025 working workbook (read-only)")
    ap.add_argument("--data-dir", type=Path, default=DEFAULT_DATA,
                    help="data directory holding item_master.csv and <year>/")
    ap.add_argument("--summary", type=Path, default=None,
                    help="also write the verification summary to this file")
    args = ap.parse_args(argv)

    if not args.inbox.is_file():
        raise SystemExit("workbook not found: %s" % args.inbox)

    summary = build(args.inbox.resolve(), args.data_dir.resolve())
    print(summary)
    if args.summary:
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text(summary + "\n", encoding="utf-8")
        print("[summary written to %s]" % args.summary)
    return 0


if __name__ == "__main__":
    sys.exit(main())
