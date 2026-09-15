"""Build the golden test-vector set in ``tests/golden/`` from a real pipeline run.

Why this exists
----------------
A JavaScript port of ``supplytrack`` needs something to check itself against
that is not "read the Python source and hope." This script runs the real,
tested Python package end to end against the same synthetic fixtures the
Python test suite uses (``tests/fixtures/*.xlsx``), with one fixed, hand-authored
set of review decisions and vendor prices, and commits every intermediate file
plus a structural dump of the finished workbook. A port that reproduces every
byte here, for the same inputs, is verified; one that does not has a concrete,
inspectable diff instead of an argument.

It also records vectors for the pure functions in isolation
(``tests/golden/unit/*.json``), so the port can be checked function-by-function
before it can run the whole pipeline.

Nothing here reads ``../inbox/`` or ``../data/``. Every input is either the
committed synthetic fixture (``tests/fixtures/amazon_2025_sample.xlsx`` /
``preferred_2025_sample.xlsx``) or a fixed decisions/prices table defined in
this file.

Determinism
-----------
Running this script twice must produce byte-identical files. The one thing
the pipeline itself is not deterministic about is wall-clock timestamps
(``ingested_at``, ``ranked_at``, the report's "report built at" cell) - these
reach the report workbook's ``Sources`` sheet and are replaced with the fixed
token ``<TIMESTAMP>`` in ``expected/report.json`` before it is written. Nothing
else in the pipeline touches the clock, the network, or random state.

Usage
-----
    cd src && python scripts/make_golden.py

Regenerate whenever supplytrack's behaviour deliberately changes, review the
diff under ``tests/golden/``, and commit it alongside the change that caused
it. ``tests/test_golden.py`` calls the same functions this script's ``main()``
does and fails if the checked-in files do not match, so a behaviour change
that forgets to regenerate the vectors is caught rather than shipped quietly.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

SRC = Path(__file__).resolve().parent.parent
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

FIXTURES_DIR = SRC / "tests" / "fixtures"
if str(FIXTURES_DIR) not in sys.path:
    sys.path.insert(0, str(FIXTURES_DIR))

import make_fixtures  # tests/fixtures/make_fixtures.py: AMAZON_ROWS, PREFERRED_ROWS

from openpyxl import load_workbook

from supplytrack import prices as prices_mod
from supplytrack import report as report_mod
from supplytrack import validate
from supplytrack.ingest import ingest, normalise_title
from supplytrack.master import (
    EXCLUDE_CATEGORIES,
    INCLUDE_CATEGORIES,
    decode_preferred_pack,
    suggest_include,
    upp_candidates,
)
from supplytrack.rank import rank
from supplytrack.xlsx import load_workbook_safe
from supplytrack.review import apply_queue, build_queue

GOLDEN = SRC / "tests" / "golden"

YEAR = 2025
TOP = 25
AMAZON_EXPORT = FIXTURES_DIR / "amazon_2025_sample.xlsx"
PREFERRED_EXPORT = FIXTURES_DIR / "preferred_2025_sample.xlsx"

# A wall-clock ISO timestamp such as "2026-09-14T21:58:50-05:00" or
# "2026-09-15T03:00:58.507029+00:00". A plain date ("2026-04-01", a typed
# checked_on value) has no "T" and does not match, so it is left alone.
_TIMESTAMP_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}")
TIMESTAMP_TOKEN = "<TIMESTAMP>"


# ---------------------------------------------------------------- fixed inputs
#
# tests/golden/inputs/decisions.csv and tests/golden/inputs/prices.csv are
# written from the constants below every time this script runs, so the
# committed input files and the logic that produced them can never drift
# apart. Both are hand-authored against the queue and ranking this script
# produces from the fixtures - see tests/golden/README.md for how each row
# was chosen.

DECISIONS_COLUMNS = [
    "key", "source", "raw_title", "amazon_category", "include",
    "units_per_pack", "canonical_name", "unit_label", "note",
]

DECISIONS_ROWS = [
    ("pbs:SAN30001", "preferred", "Permanent Marker Fine Point Black", "",
     "y", "1", "Permanent Marker Fine Point Black", "EA", ""),
    ("amz:bic round stic ballpoint pens medium point black 60-count", "amazon",
     "BIC Round Stic Ballpoint Pens, Medium Point, Black, 60-Count", "Office Product",
     "y", "60", "BIC Ballpoint Pens Black Medium", "EA", ""),
    ("amz:bubly sparkling water variety pack 18 count", "amazon",
     "Bubly Sparkling Water, Variety Pack, 18 Count", "Grocery",
     "n", "", "Bubly Sparkling Water", "EA", "excluded: Grocery item is not an office supply"),
    ("amz:post-it notes 3 x 3 inches canary yellow 12 pads", "amazon",
     "Post-it Notes, 3 x 3 Inches, Canary Yellow, 12 Pads", "Office Product",
     "y", "12", "Post-it Notes 3x3 Yellow", "EA", ""),
    ("pbs:PBS9001", "preferred", "Letterhead Printed 2 Colour", "",
     "y", "1", "Letterhead Printed 2 Colour", "EA", ""),
    ("pbs:UNV35668", "preferred", "Binder Clips Medium Black", "",
     "y", "100", "Binder Clips Medium Black", "EA", ""),
    ("pbs:UNV21200", "preferred", "Copy Paper 8.5 x 11 20lb White", "",
     "y", "10", "Preferred Copy Paper 8.5x11 White", "EA", ""),
    ("amz:hammermill copy paper 8.5 x 11 500 sheets/ream 10 reams/carton", "amazon",
     "Hammermill Copy Paper, 8.5 x 11, 500 Sheets/Ream, 10 Reams/Carton", "Office Product",
     "y", "10", "Hammermill Copy Paper 8.5x11", "RM",
     "purchase unit is the carton: 10 reams per carton"),
    ("amz:kleenex professional facial tissue display pack of 6", "amazon",
     "Kleenex Professional Facial Tissue, Display Pack of 6",
     "Business, Industrial, & Scientific Supplies Basic",
     "y", "6", "Kleenex Professional Facial Tissue", "EA", ""),
    ("amz:avery printable name tags 400 printable name tag inserts 3 x 4 inches", "amazon",
     "Avery Printable Name Tags, 400 Printable Name Tag Inserts, 3 x 4 Inches", "Office Product",
     "y", "400", "Avery Name Tag Inserts", "EA", ""),
    ("amz:expo low-odor dry erase markers chisel tip assorted 48/pack", "amazon",
     "EXPO Low-Odor Dry Erase Markers, Chisel Tip, Assorted, 48/Pack", "Office Product",
     "y", "48", "EXPO Dry Erase Markers Assorted", "EA", ""),
    # Deliberately differs from the regex suggestion (12, from "12 Count"):
    # a person confirmed the box actually holds two dozen. Also exercises
    # validate.UPP_TITLE_MISMATCH, since 24 neither appears in the title nor
    # is a product of its candidates.
    ("amz:sharpie permanent markers fine point assorted 12 count", "amazon",
     "Sharpie Permanent Markers, Fine Point, Assorted, 12 Count", "Office Product",
     "y", "24", "Sharpie Permanent Markers Assorted", "EA",
     "title says 12 Count; box actually holds two dozen; confirmed against invoice"),
    # Merges with the Avery row above under one canonical_name, and its
    # units_per_pack (200) deliberately differs from the regex suggestion
    # (400, same as the title states) - the merge-with-disagreeing-pack-sizes
    # case, so ranked.csv's units_per_pack comes back blank for this item.
    ("amz:avery name badge inserts 400 printable name tag inserts 3 x 4 in white", "amazon",
     "Avery Name Badge Inserts, 400 Printable Name Tag Inserts, 3 x 4 in, White", "Office Product",
     "y", "200", "Avery Name Tag Inserts", "EA",
     "listing states 400 but this SKU ships a half case (200); merged with the other "
     "Name Tag listing"),
    ("amz:cable zip ties 8 inch black (bulk 1000 pack)", "amazon",
     "Cable Zip Ties, 8 Inch, Black (Bulk 1000 Pack)", "Tools & Home Improvement",
     "n", "", "Cable Zip Ties", "EA", "excluded by the office manager in 2025"),
    ("amz:scotch magic tape 3/4 x 1000 inches 6 rolls", "amazon",
     "Scotch Magic Tape, 3/4 x 1000 Inches, 6 Rolls", "Office Product",
     "y", "6", "Scotch Magic Tape 3/4in", "EA", ""),
]

PRICES_COLUMNS = ["rank", "canonical_name", "vendor", "unit_price", "status", "url", "checked_on", "note"]

# One row per (item x vendor) for the 12 items the decisions above produce,
# in rank order. Deliberately includes every status: priced (most cells),
# not_available (BIC/Staples, Letterhead/Staples), discontinued
# (Post-it/Staples) and unpriced (Preferred Copy Paper/Amazon, left blank -
# nobody has looked it up). One unit_price ("$9.50") carries a leading "$" to
# exercise the same tolerant parsing the real prices.csv needs.
PRICES_ROWS = [
    ("1", "Avery Name Tag Inserts", "Office Depot", "3.10", "priced", "https://example.com/avery-od", "2026-04-01", ""),
    ("1", "Avery Name Tag Inserts", "Preferred", "2.95", "priced", "https://example.com/avery-pref", "2026-04-01", ""),
    ("1", "Avery Name Tag Inserts", "Amazon", "3.05", "priced", "https://example.com/avery-amz", "2026-04-01", "last paid per each: 0.1105"),
    ("1", "Avery Name Tag Inserts", "Staples", "3.20", "priced", "https://example.com/avery-stpl", "2026-04-01", ""),
    ("2", "BIC Ballpoint Pens Black Medium", "Office Depot", "0.22", "priced", "https://example.com/bic-od", "2026-04-01", ""),
    ("2", "BIC Ballpoint Pens Black Medium", "Preferred", "0.19", "priced", "https://example.com/bic-pref", "2026-04-01", ""),
    ("2", "BIC Ballpoint Pens Black Medium", "Amazon", "0.20", "priced", "https://example.com/bic-amz", "2026-04-01", "last paid per each: 0.0997"),
    ("2", "BIC Ballpoint Pens Black Medium", "Staples", "", "not_available", "", "2026-04-01", "discontinued line at Staples"),
    ("3", "Binder Clips Medium Black", "Office Depot", "1.05", "priced", "https://example.com/binder-od", "2026-04-01", ""),
    ("3", "Binder Clips Medium Black", "Preferred", "0.95", "priced", "https://example.com/binder-pref", "2026-04-01", ""),
    ("3", "Binder Clips Medium Black", "Amazon", "1.00", "priced", "https://example.com/binder-amz", "2026-04-01", ""),
    ("3", "Binder Clips Medium Black", "Staples", "1.10", "priced", "https://example.com/binder-stpl", "2026-04-01", ""),
    ("4", "EXPO Dry Erase Markers Assorted", "Office Depot", "$9.50", "priced", "https://example.com/expo-od", "2026-04-01", ""),
    ("4", "EXPO Dry Erase Markers Assorted", "Preferred", "8.60", "priced", "https://example.com/expo-pref", "2026-04-01", ""),
    ("4", "EXPO Dry Erase Markers Assorted", "Amazon", "8.90", "priced", "https://example.com/expo-amz", "2026-04-01", "last paid per each: 0.8331"),
    ("4", "EXPO Dry Erase Markers Assorted", "Staples", "9.75", "priced", "https://example.com/expo-stpl", "2026-04-01", ""),
    ("5", "Post-it Notes 3x3 Yellow", "Office Depot", "4.75", "priced", "https://example.com/postit-od", "2026-04-01", ""),
    ("5", "Post-it Notes 3x3 Yellow", "Preferred", "4.30", "priced", "https://example.com/postit-pref", "2026-04-01", ""),
    ("5", "Post-it Notes 3x3 Yellow", "Amazon", "4.45", "priced", "https://example.com/postit-amz", "2026-04-01", "last paid per each: 1.2292"),
    ("5", "Post-it Notes 3x3 Yellow", "Staples", "", "discontinued", "", "2026-04-02", ""),
    ("6", "Sharpie Permanent Markers Assorted", "Office Depot", "0.65", "priced", "https://example.com/sharpie-od", "2026-04-01", ""),
    ("6", "Sharpie Permanent Markers Assorted", "Preferred", "0.58", "priced", "https://example.com/sharpie-pref", "2026-04-01", ""),
    ("6", "Sharpie Permanent Markers Assorted", "Amazon", "0.60", "priced", "https://example.com/sharpie-amz", "2026-04-01", "last paid per each: 0.4717"),
    ("6", "Sharpie Permanent Markers Assorted", "Staples", "0.68", "priced", "https://example.com/sharpie-stpl", "2026-04-01", ""),
    ("7", "Hammermill Copy Paper 8.5x11", "Office Depot", "44.25", "priced", "https://example.com/hammermill-od", "2026-04-01", ""),
    ("7", "Hammermill Copy Paper 8.5x11", "Preferred", "39.99", "priced", "https://example.com/hammermill-pref", "2026-04-01", ""),
    ("7", "Hammermill Copy Paper 8.5x11", "Amazon", "41.10", "priced", "https://example.com/hammermill-amz", "2026-04-01", "last paid per each: 5.499"),
    ("7", "Hammermill Copy Paper 8.5x11", "Staples", "46.00", "priced", "https://example.com/hammermill-stpl", "2026-04-01", ""),
    ("8", "Preferred Copy Paper 8.5x11 White", "Office Depot", "4.30", "priced", "https://example.com/prefpaper-od", "2026-04-01", ""),
    ("8", "Preferred Copy Paper 8.5x11 White", "Preferred", "3.95", "priced", "https://example.com/prefpaper-pref", "2026-04-01", ""),
    ("8", "Preferred Copy Paper 8.5x11 White", "Amazon", "", "unpriced", "", "", ""),
    ("8", "Preferred Copy Paper 8.5x11 White", "Staples", "4.55", "priced", "https://example.com/prefpaper-stpl", "2026-04-01", ""),
    ("9", "Kleenex Professional Facial Tissue", "Office Depot", "12.49", "priced", "https://example.com/kleenex-od", "2026-04-01", ""),
    ("9", "Kleenex Professional Facial Tissue", "Preferred", "11.20", "priced", "https://example.com/kleenex-pref", "2026-04-01", ""),
    ("9", "Kleenex Professional Facial Tissue", "Amazon", "11.60", "priced", "https://example.com/kleenex-amz", "2026-04-01", "last paid per each: 3.1"),
    ("9", "Kleenex Professional Facial Tissue", "Staples", "12.95", "priced", "https://example.com/kleenex-stpl", "2026-04-01", ""),
    ("10", "Permanent Marker Fine Point Black", "Office Depot", "1.10", "priced", "https://example.com/marker-od", "2026-04-01", ""),
    ("10", "Permanent Marker Fine Point Black", "Preferred", "0.98", "priced", "https://example.com/marker-pref", "2026-04-01", ""),
    ("10", "Permanent Marker Fine Point Black", "Amazon", "1.02", "priced", "https://example.com/marker-amz", "2026-04-01", ""),
    ("10", "Permanent Marker Fine Point Black", "Staples", "1.15", "priced", "https://example.com/marker-stpl", "2026-04-01", ""),
    ("11", "Scotch Magic Tape 3/4in", "Office Depot", "5.49", "priced", "https://example.com/scotch-od", "2026-04-01", ""),
    ("11", "Scotch Magic Tape 3/4in", "Preferred", "4.95", "priced", "https://example.com/scotch-pref", "2026-04-01", ""),
    ("11", "Scotch Magic Tape 3/4in", "Amazon", "5.10", "priced", "https://example.com/scotch-amz", "2026-04-01", "last paid per each: 2.165"),
    ("11", "Scotch Magic Tape 3/4in", "Staples", "5.75", "priced", "https://example.com/scotch-stpl", "2026-04-01", ""),
    ("12", "Letterhead Printed 2 Colour", "Office Depot", "0.42", "priced", "https://example.com/letterhead-od", "2026-04-01", ""),
    ("12", "Letterhead Printed 2 Colour", "Preferred", "0.38", "priced", "https://example.com/letterhead-pref", "2026-04-01", ""),
    ("12", "Letterhead Printed 2 Colour", "Amazon", "0.40", "priced", "https://example.com/letterhead-amz", "2026-04-01", ""),
    ("12", "Letterhead Printed 2 Colour", "Staples", "", "not_available", "", "2026-04-01", ""),
]


def _write_csv(path: Path, columns: list[str], rows: list[tuple]) -> bytes:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(columns)
        writer.writerows(rows)
    return path.read_bytes()


def _write_inputs(inputs_dir: Path) -> dict[str, bytes]:
    return {
        "inputs/decisions.csv": _write_csv(inputs_dir / "decisions.csv", DECISIONS_COLUMNS, DECISIONS_ROWS),
        "inputs/prices.csv": _write_csv(inputs_dir / "prices.csv", PRICES_COLUMNS, PRICES_ROWS),
        "inputs/fixture_tables.json": _json_bytes(_fixture_tables()),
    }


# ------------------------------------------------------------ fixture tables

# The workbooks sheet recognition has to cope with: the working workbook as the
# office manager keeps it, the same after Amazon changed the columns, and three
# files that must be refused with a message that says why.
RECOGNITION_FIXTURES = {
    "working": "working_workbook_2025_sample.xlsx",
    "filtered": "working_workbook_filtered_2025_sample.xlsx",
    "variant": "working_workbook_variant_2025_sample.xlsx",
    "refunds": "amazon_refunds_sample.xlsx",
    "two_amazon": "two_amazon_sheets_sample.xlsx",
    "missing_column": "amazon_missing_column_sample.xlsx",
}


def _sheet_rows(path: Path) -> list[list]:
    """Every cell of the first sheet, exactly as openpyxl hands it over.

    Deliberately *not* normalised: no header case-folding, no blank-row
    dropping, no padding. This is the raw shape a browser gets from ExcelJS
    after it opens the same workbook, and the port's own `normaliseTable` is
    what has to turn it into a table - so normalising here would skip the very
    step the ingest parity test exists to check.

    Dates are the one conversion. openpyxl returns a ``datetime`` for a
    date-formatted cell, which has no JSON form; ISO 8601 is the shape the
    port's ``coerceDate`` already accepts alongside a real Date and an Excel
    serial number.
    """
    # The same tolerant reader ingest uses, so a fixture carrying the Amazon
    # export's out-of-range font family still opens here.
    wb = load_workbook_safe(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return [
            [v.isoformat() if isinstance(v, (dt.datetime, dt.date)) else v for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
    finally:
        wb.close()


def _all_sheets(path: Path) -> list[dict]:
    """Every sheet of a workbook as ``{name, grid}``, dates ISO-ified.

    The same conversion ``_sheet_rows`` makes, for the same reason: a datetime
    has no JSON form, and ISO 8601 is a shape the port's ``coerceDate`` and its
    header scan both already handle.
    """
    wb = load_workbook_safe(path, data_only=True)
    try:
        return [
            {
                "name": name,
                "grid": [
                    [v.isoformat() if isinstance(v, (dt.datetime, dt.date)) else v for v in row]
                    for row in wb[name].iter_rows(values_only=True)
                ],
            }
            for name in wb.sheetnames
        ]
    finally:
        wb.close()


def _workbook_entry(path: Path) -> dict:
    return {"file": path.name, "sheets": _all_sheets(path)}


def _fixture_tables() -> dict:
    """The fixture exports as plain tables, so a Node test can run ingest.

    The fixtures are ``.xlsx`` and reading a workbook is the browser's job, not
    the parity suite's - depending on ExcelJS there would mean the suite no
    longer runs on Node alone. Dumping the sheets here lets the port's ``ingest``
    be checked against ``expected/lines.csv`` with no spreadsheet reader at all.

    ``amazon`` and ``preferred`` are the first sheet of each standalone export
    and are what the end-to-end vector runs on. ``files`` is the same pair in
    the shape the page hands over, one entry per file with every sheet named,
    and ``workbooks`` adds the awkward files sheet recognition has to cope with.
    """
    return {
        "amazon": _sheet_rows(AMAZON_EXPORT),
        "preferred": _sheet_rows(PREFERRED_EXPORT),
        "amazon_file": AMAZON_EXPORT.name,
        "preferred_file": PREFERRED_EXPORT.name,
        "files": [_workbook_entry(AMAZON_EXPORT), _workbook_entry(PREFERRED_EXPORT)],
        "workbooks": {
            name: _workbook_entry(FIXTURES_DIR / filename)
            for name, filename in RECOGNITION_FIXTURES.items()
        },
    }


# --------------------------------------------------------------------- json


def _json_bytes(obj) -> bytes:
    return (json.dumps(obj, indent=2, ensure_ascii=False, sort_keys=True) + "\n").encode("utf-8")


def _normalise_scalar(value):
    if isinstance(value, str) and _TIMESTAMP_RE.match(value):
        return TIMESTAMP_TOKEN
    return value


def _normalise_deep(value):
    """``_normalise_scalar`` through nested dicts and lists.

    ``run.json`` carries its wall-clock ``ingested_at`` at the top level, but a
    future field could nest one inside ``warnings`` or ``non_closed``; walking
    the whole structure means a new timestamp cannot quietly make the vectors
    non-deterministic.
    """
    if isinstance(value, dict):
        return {k: _normalise_deep(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_normalise_deep(v) for v in value]
    return _normalise_scalar(value)


def _color_payload(color):
    if color is None:
        return None
    if color.type == "rgb":
        return {"type": "rgb", "rgb": color.rgb}
    if color.type == "theme":
        return {"type": "theme", "theme": color.theme, "tint": color.tint}
    if color.type == "indexed":
        return {"type": "indexed", "indexed": color.indexed}
    return {"type": color.type}


def _fill_payload(fill) -> dict:
    return {
        "patternType": fill.patternType,
        "fgColor": _color_payload(fill.fgColor),
        "bgColor": _color_payload(fill.bgColor),
    }


def dump_workbook(path: Path) -> dict:
    """Every non-empty cell, merges, column widths, freeze panes, conditional
    formatting and non-default fills of every sheet, as a plain JSON-safe dict.

    Timestamps are normalised (see ``_normalise_scalar``); everything else is
    exactly what a JS port that opens the same workbook should see.
    """
    wb = load_workbook(path, data_only=False)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        cells: dict[str, dict] = {}
        fills: dict[str, dict] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.fill is not None and cell.fill.patternType:
                    fills[cell.coordinate] = _fill_payload(cell.fill)
                if cell.value is not None:
                    is_formula = cell.data_type == "f"
                    value = _normalise_scalar(cell.value)
                    cells[cell.coordinate] = {
                        "v": value,
                        "f": value if is_formula else None,
                        "t": cell.data_type,
                    }

        cf_rules = []
        for rng in ws.conditional_formatting:
            sqref = str(rng.sqref)
            for rule in ws.conditional_formatting[rng]:
                highlight_fill = None
                if rule.dxf is not None and rule.dxf.fill is not None:
                    highlight_fill = _fill_payload(rule.dxf.fill)
                cf_rules.append(
                    {
                        "sqref": sqref,
                        "type": rule.type,
                        "formula": list(rule.formula or []),
                        "highlight_fill": highlight_fill,
                    }
                )
        cf_rules.sort(key=lambda r: r["sqref"])

        sheets[name] = {
            "cells": cells,
            "fills": fills,
            "merged_cells": sorted(str(r) for r in ws.merged_cells.ranges),
            "column_widths": {
                letter: dim.width for letter, dim in ws.column_dimensions.items() if dim.width is not None
            },
            "freeze_panes": ws.freeze_panes,
            "conditional_formatting": cf_rules,
        }
    return {"sheet_order": list(wb.sheetnames), "sheets": sheets}


# ---------------------------------------------------------------- pipeline run


def _run_pipeline(work_dir: Path, inputs_dir: Path) -> dict[str, bytes]:
    data_dir = work_dir / "data"
    out_dir = work_dir / "out"
    data_dir.mkdir(parents=True)

    ingest(data_dir, YEAR, AMAZON_EXPORT, PREFERRED_EXPORT)
    lines_bytes = (data_dir / str(YEAR) / "lines.csv").read_bytes()

    # Captured here, before rank appends ranked_at and its warnings: this is
    # ingest's own output, and it is the only place `rows_written` and the
    # non-Closed order statuses exist. Neither can be recovered from lines.csv,
    # so without this the port cannot reproduce LINE_COUNT_MISMATCH or
    # STATUS_NOT_CLOSED at all.
    run_payload = _normalise_deep(
        json.loads((data_dir / str(YEAR) / "run.json").read_text(encoding="utf-8"))
    )

    # Captured before apply: this is the queue a reviewer actually sees.
    queue_path = build_queue(data_dir, YEAR)
    queue_bytes = queue_path.read_bytes()

    apply_queue(data_dir, YEAR, inputs_dir / "decisions.csv", proposed=False)
    master_bytes = (data_dir / "item_master.csv").read_bytes()

    rank(data_dir, YEAR)
    ranked_bytes = (data_dir / str(YEAR) / "ranked.csv").read_bytes()
    excluded_bytes = (data_dir / str(YEAR) / "excluded.csv").read_bytes()

    # Captured as the template, before the fixed filled sheet is swapped in.
    prices_mod.write_template(data_dir, YEAR, TOP)
    prices_template_bytes = (data_dir / str(YEAR) / "prices.csv").read_bytes()

    shutil.copyfile(inputs_dir / "prices.csv", data_dir / str(YEAR) / "prices.csv")

    findings = validate.run_all(data_dir, YEAR, TOP)
    findings_payload = [{"level": f.level, "code": f.code, "message": f.message} for f in findings]

    report_path = report_mod.build_report(data_dir, out_dir, YEAR, TOP)
    report_payload = dump_workbook(report_path)

    return {
        "inputs/run.json": _json_bytes(run_payload),
        "expected/lines.csv": lines_bytes,
        "expected/review_queue.csv": queue_bytes,
        "expected/item_master.csv": master_bytes,
        "expected/ranked.csv": ranked_bytes,
        "expected/excluded.csv": excluded_bytes,
        "expected/prices.csv": prices_template_bytes,
        "expected/findings.json": _json_bytes(findings_payload),
        "expected/report.json": _json_bytes(report_payload),
    }


# ------------------------------------------------------------- unit vectors

NORMALISE_TITLE_INPUTS = [
    "  BIC  Round   Stic  ",
    "3/4 x 1000 (Bulk) - 8.5",
    "Sharpie®, Fine Point; #2",
    "Café Notebook",
    "",
    None,
    "ALL CAPS TITLE",
    "already lowercase title",
    "Tabs\tand\nnewlines\tcollapse",
    "Product™ Name©",
    "Product #123 (Model-X)",
    "50% Off Sale!!!",
    "Multi   Space    Title",
    "Ünïcödé Áccénts",  # unicode accents, plain-stripped
    "Ｆｕｌｌｗｉｄｔｈ",  # fullwidth "Fullwidth" -> NFKC folds to ASCII
    "ﬁle folder ﬂap",  # ligatures fi/fl -> NFKC decomposes to "fi"/"fl"
    "América Brand",  # combining acute accent, composes under NFKC then is stripped
    "Product—Name Dash",  # em dash, distinct from the kept ascii hyphen
    "Kirkland’s “Best” Paper",  # smart quotes
    "8.5 x 11 Paper",
    "N/A Widget",
    "Item#42: Deluxe (Pro), v2.0!",
    "paper clips",
    "!!!@@@###",
    "   ",
    ".Test.",
    "Price: €5.00 Each",
    "((Test))",
    "Foo Bar",  # non-breaking space
    "咖啡 Notebook",  # CJK characters
    "CO₂ Absorber",  # subscript 2 -> NFKC folds to ascii "2"
    "１２３ Widget",  # fullwidth digits -> NFKC folds to "123"
]

UPP_CANDIDATES_INVENTED_INPUTS = [
    "Sturdy 2 Inch Ring Binders, Black",
    "Heavy Duty 3 Ring Binder, Blue",
    "Index Dividers, 5 Tab, Multicolour",
    "Index Dividers, 26 Tab, A-Z",
    'Presentation Binder, 1" Capacity',
    "Laminating Pouches, 5 mil, Letter Size",
    "Paper Towels, 2 Ply, White",
    "Copier Paper, 20 lb, Bright White",
    "Filing Envelopes 9 x 12, Kraft",
    "Poster Board 22 x 28, Assorted",
    "Ink Cartridge PFI1000 Matte Black",
    "Copy Paper (38111), White",
    "Pack of 6 Highlighters",
    "Pack of 24 Markers",
    "12-Count Pencils",
    "500-Count Sticky Notes",
    "48/Pack Dry Erase Markers",
    "6/Pack Correction Tape",
    "Case of 10 Packs Legal Pads",
    "Case of 4 Packs Paper Towels",
    "Toner 3.5 Pack",  # decimal: the 5 must not be read as a pack size of its own
    "Legal Pads, 8.5 x 14, 50 Sheets",
    "3 x 5 Index Cards, 100 Count",
    "100 Pack",
    "1000-Pack",
    "Binder Clips, Box of 50",
    "Sticky Notes, Bundle of 12",
    "Tape Dispenser, Set of 3",
    "Paper, 10 Reams per Carton",
    "Report Covers, 25 Sheets per Box",
    "Envelopes, 100/Box",
    "Two Dozen Pencils",  # a word, not a digit: no candidate
    "12 Dozen Pencils",  # "dozen" is not a counted phrase in upp_candidates
    "9 x 12 Envelopes, Pack of 250",
    "Notebook, 100 Pages",  # "pages" is a measurement word
    "5 Star Rated Stapler",
    "Binder 3-Ring 2 Inch White",
    "12/Pack, Case of 10 Packs",
    "1000 Sheets/Ream, 5 Reams/Carton",
    "Highlighters, 6ct, Assorted Colors",
    "Sticky Flags, 4 Pads of 50 Sheets",
]

DECODE_PACK_INPUTS = [
    "CT10", "BX100", "CS1", "Each", "EA", "ea.", "1", "CT2500", "pk12",
    "", "   ", "ASSORTED", "BX-", None, "CT0", "0", "cs1", "XYZ1234567",
    "each ", " Each", "PK5",
]


def _suggest_include_inputs() -> list[tuple[str | None, str]]:
    categories = (
        sorted(INCLUDE_CATEGORIES)
        + sorted(EXCLUDE_CATEGORIES)
        + ["Tools & Home Improvement", "", "  Office Product  ", "GROCERY", None]
    )
    return [(category, source) for category in categories for source in ("amazon", "preferred")]


def _fixture_titles() -> list[str]:
    amazon_titles = [row[4] for row in make_fixtures.AMAZON_ROWS]
    preferred_titles = [row[2] for row in make_fixtures.PREFERRED_ROWS]
    return amazon_titles + preferred_titles


def _build_unit_vectors() -> dict[str, bytes]:
    normalise_title_vectors = [{"input": x, "output": normalise_title(x)} for x in NORMALISE_TITLE_INPUTS]

    upp_inputs = _fixture_titles() + UPP_CANDIDATES_INVENTED_INPUTS
    upp_vectors = [
        {"input": title, "output": [[value, phrase] for value, phrase in upp_candidates(title)]}
        for title in upp_inputs
    ]

    decode_vectors = [{"input": code, "output": decode_preferred_pack(code)} for code in DECODE_PACK_INPUTS]

    suggest_vectors = []
    for category, source in _suggest_include_inputs():
        include, reason = suggest_include(category, source)
        suggest_vectors.append(
            {
                "input": {"amazon_category": category, "source": source},
                "output": {"include": include, "reason": reason},
            }
        )

    return {
        "unit/normalise_title.json": _json_bytes(normalise_title_vectors),
        "unit/upp_candidates.json": _json_bytes(upp_vectors),
        "unit/decode_preferred_pack.json": _json_bytes(decode_vectors),
        "unit/suggest_include.json": _json_bytes(suggest_vectors),
    }


# --------------------------------------------------------------------- driver


def build_everything(work_dir: Path) -> dict[str, bytes]:
    """Compute every golden file's bytes. Pure with respect to the filesystem
    outside ``work_dir``: nothing under ``tests/golden/`` is touched here."""
    work_dir = Path(work_dir)
    inputs_dir = work_dir / "inputs"
    files: dict[str, bytes] = {}
    files.update(_write_inputs(inputs_dir))
    files.update(_run_pipeline(work_dir / "pipeline", inputs_dir))
    files.update(_build_unit_vectors())
    return files


def write_golden(files: dict[str, bytes]) -> None:
    for rel_path, content in files.items():
        dest = GOLDEN / rel_path
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(content)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="supplytrack-golden-") as tmp:
        files = build_everything(Path(tmp))
    write_golden(files)
    print(f"Wrote {len(files)} file(s) under {GOLDEN}")
    for rel_path in sorted(files):
        print(f"  {rel_path}  ({len(files[rel_path])} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
