"""Tests for supplytrack.report: the leadership workbook builder."""
from __future__ import annotations

import csv
import json
import shutil
import sys
import types
from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import load_workbook


def _ensure_fallback_modules() -> None:
    """Stand in for errors.py / validate.py if they are not built yet.

    See the identical helper in test_prices.py for the rationale: another
    agent is building the rest of the package concurrently, and this file
    should not fail to import merely because a sibling module isn't there
    yet.
    """
    try:
        import supplytrack.errors  # noqa: F401
    except ImportError:
        mod = types.ModuleType("supplytrack.errors")

        class SupplytrackError(Exception):
            pass

        mod.SupplytrackError = SupplytrackError
        sys.modules["supplytrack.errors"] = mod

    try:
        import supplytrack.validate  # noqa: F401
    except ImportError:
        mod = types.ModuleType("supplytrack.validate")

        @dataclass
        class Finding:
            level: str
            code: str
            message: str

        mod.Finding = Finding
        sys.modules["supplytrack.validate"] = mod


_ensure_fallback_modules()

from supplytrack.report import build_report  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "report"


def _make_data_dir(tmp_path: Path, year: int = 2025) -> Path:
    """<tmp>/data/<year>/{ranked,prices}.csv, copied from the fixtures."""
    data_dir = tmp_path / "data"
    year_dir = data_dir / str(year)
    year_dir.mkdir(parents=True)
    shutil.copy(FIXTURES / "ranked.csv", year_dir / "ranked.csv")
    shutil.copy(FIXTURES / "prices.csv", year_dir / "prices.csv")
    return data_dir


def _read_ranked_rows() -> list[dict[str, str]]:
    with (FIXTURES / "ranked.csv").open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


# --- full layout, all optional sheets present -----------------------------


@pytest.fixture
def full_data_dir(tmp_path) -> Path:
    data_dir = _make_data_dir(tmp_path, 2025)
    year_dir = data_dir / "2025"

    (year_dir / "run.json").write_text(
        json.dumps(
            {
                "year": 2025,
                "amazon_file": "amazon-2025.xlsx",
                "preferred_file": "preferred-2025.xlsx",
                "amazon_rows": 1212,
                "preferred_rows": 26,
                "date_min": "2025-01-02",
                "date_max": "2025-12-30",
                "ingested_at": "2026-03-01T00:00:00+00:00",
                "warnings": ["3 lines outside the reporting year"],
            }
        ),
        encoding="utf-8",
    )

    with (year_dir / "excluded.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["key", "source", "raw_title", "amazon_category", "packs", "reason"])
        writer.writerow(["amz:zip-ties", "amazon", "ARTISTRO Zip Ties", "Grocery", "5", "include=n"])
        writer.writerow(["amz:unknown-thing", "amazon", "Mystery Widget", "Office Product", "2", "not in item master"])

    ranked_rows = _read_ranked_rows()
    with (data_dir / "item_master.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "key", "source", "raw_title", "include", "canonical_name", "units_per_pack",
            "unit_label", "upp_source", "amazon_category", "first_seen", "last_seen", "note",
        ])
        for row in ranked_rows:
            writer.writerow([
                row["keys"], row["sources"], row["canonical_name"], "y", row["canonical_name"],
                row["units_per_pack"], row["unit_label"], row["upp_source"], "Office Product",
                "2025", "2025", "",
            ])

    return data_dir


def test_build_report_writes_expected_path_and_sheet_names(tmp_path, full_data_dir):
    out_dir = tmp_path / "out"
    path = build_report(full_data_dir, out_dir, 2025, 25)

    assert path == out_dir / "Acme Widget Top 25 Items Comparison 2025.xlsx"
    assert path.exists()

    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["Top 25", "All items", "Excluded", "Sources"]


def test_top_sheet_title_headers_and_data_rows(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Top 25"]

    assert ws["B4"].value == "Top 25 Items By Quantity"
    assert ws["B4"].font.bold is True
    assert ws["B4"].font.size == 14
    assert "B4:N4" in ws.merged_cells

    assert ws["F5"].value == "Unit Price (per each)"
    assert "F5:I5" in ws.merged_cells
    assert ws["K5"].value == "Total Cost at 2025 volume"
    assert "K5:N5" in ws.merged_cells

    assert [ws[f"{c}6"].value for c in "BCDE"] == ["Ranking", "Description", "Pack/Size", "Quantity"]
    assert [ws[f"{c}6"].value for c in "FGHI"] == ["Office Depot", "Preferred", "Amazon", "Staples"]
    assert [ws[f"{c}6"].value for c in "KLMN"] == ["Office Depot", "Preferred", "Amazon", "Staples"]

    # Row 7 = rank 1, "Copy Paper 8.5x11 White", unit_label RM, eaches 9600.
    assert ws["B7"].value == 1
    assert ws["C7"].value == "Copy Paper 8.5x11 White"
    assert ws["D7"].value == "1 RM"
    assert ws["E7"].value == 9600
    assert ws["E7"].number_format == "#,##0"
    assert ws["F7"].number_format == "$#,##0.00"
    assert ws["K7"].value == "=F7*E7"
    assert ws["N7"].value == "=I7*E7"


def test_text_statuses_land_as_text_not_available_and_discontinued(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Top 25"]

    # Fixture: rank 5 (row 11) Preferred = not_available; rank 12 (row 18) Staples = discontinued;
    # rank 9 (row 15) Amazon = unpriced.
    assert ws["G11"].value == "not available"
    assert ws["I18"].value == "Discontinued"
    assert ws["H15"].value == "not priced"
    # The paired total-cost cell carries the same text, not a formula.
    assert ws["L11"].value == "not available"
    assert ws["N18"].value == "Discontinued"
    assert ws["M15"].value == "not priced"


def test_totals_priced_items_and_comparable_basket_rows(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Top 25"]

    last = 31  # 6 + top(25)
    label_row, totals_row, priced_row, basket_row = 32, 33, 34, 35

    assert ws[f"K{label_row}"].value == "Office Depot Total Cost"
    assert ws[f"N{label_row}"].value == "Staples Total Cost"

    assert ws[f"K{totals_row}"].value == f"=SUM(K7:K{last})"
    assert ws[f"N{totals_row}"].value == f"=SUM(N7:N{last})"
    # No sum of unit prices.
    assert ws[f"F{totals_row}"].value is None

    assert ws[f"F{priced_row}"].value == f'=COUNT(F7:F{last})&" / 25"'
    assert ws[f"I{priced_row}"].value == f'=COUNT(I7:I{last})&" / 25"'

    basket_formula = ws[f"K{basket_row}"].value
    assert basket_formula.startswith("=SUM(")
    # Rank 5 (row 11, Preferred not_available), rank 12 (row 18, Staples
    # discontinued) and rank 9 (row 15, Amazon unpriced) do not qualify for
    # the comparable basket - a text cell in any vendor's column disqualifies
    # the whole row, in every total-cost column, not just the vendor with
    # the text cell.
    assert "K11" not in basket_formula
    assert "K18" not in basket_formula
    assert "K15" not in basket_formula
    assert "K7" in basket_formula

    legend_row = basket_row + 2
    assert ws[f"B{legend_row}"].value == "Lowest price for row item"
    assert ws[f"B{legend_row}"].fill.fgColor.rgb.endswith("FFF2CC")


def test_conditional_formatting_rules(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Top 25"]

    cf = ws.conditional_formatting
    ranges = {rng.sqref: cf[rng] for rng in cf}
    assert set(ranges.keys()) == {"F7:I31", "K7:N31"}
    assert sum(len(rules) for rules in ranges.values()) == 2

    price_rule = ranges["F7:I31"][0]
    assert price_rule.formula == ["AND(ISNUMBER(F7),F7=MIN($F7:$I7))"]
    assert price_rule.dxf.fill.bgColor.rgb.endswith("FFF2CC")

    total_rule = ranges["K7:N31"][0]
    assert total_rule.formula == ["AND(ISNUMBER(K7),K7=MIN($K7:$N7))"]


def test_column_widths_freeze_panes_and_wrap_text(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Top 25"]

    assert ws.column_dimensions["B"].width == 9
    assert ws.column_dimensions["C"].width == 60
    assert ws.column_dimensions["D"].width == 10
    assert ws.column_dimensions["E"].width == 11
    for col in "FGHI":
        assert ws.column_dimensions[col].width == 13
    assert ws.column_dimensions["J"].width == 2
    for col in "KLMN":
        assert ws.column_dimensions[col].width == 15

    assert ws.freeze_panes == "C7"
    assert ws["C7"].alignment.wrap_text is True

    # Column J divider fill spans the header and every data row.
    assert ws["J6"].fill.fgColor.rgb.endswith("FFF2CC")
    assert ws["J31"].fill.fgColor.rgb.endswith("FFF2CC")


def test_all_items_sheet_full_ranked_csv(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["All items"]

    assert ws["A1"].value == "rank"
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:J31"  # 10 columns, 30 data rows + header

    # All 30 invented items present, not just the top 25.
    names = [ws.cell(row=r, column=2).value for r in range(2, 32)]
    assert len(names) == 30
    assert "Binder 3-Ring 2 Inch White" in names  # rank 30, outside the top 25


def test_excluded_sheet_present_when_file_exists(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Excluded"]

    assert ws["A1"].value == "key"
    assert ws["A1"].font.bold is True
    assert ws.cell(row=2, column=1).value == "amz:zip-ties"
    assert ws.cell(row=3, column=1).value == "amz:unknown-thing"


def test_sources_sheet_has_run_json_keys_and_derived_rows(tmp_path, full_data_dir):
    path = build_report(full_data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)
    ws = wb["Sources"]

    kv = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}

    assert kv["year"] == 2025
    assert kv["amazon_file"] == "amazon-2025.xlsx"
    assert kv["amazon_rows"] == 1212
    assert kv["warnings"] == "3 lines outside the reporting year"
    assert kv["item master rows"] == 30
    assert kv["top_n"] == 25
    assert kv["unpriced cells"] == 1  # rank 9 (row 15), Amazon, in the fixture
    assert "report built at" in kv
    # ISO timestamp, parseable.
    from datetime import datetime

    datetime.fromisoformat(kv["report built at"])

    non_master = kv["items with upp_source != master"]
    for name in ("Sticky Notes 3x3 Yellow", "Rubber Bands Assorted Size", "Colored Copy Paper Assorted"):
        assert name in non_master
    assert "Copy Paper 8.5x11 White" not in non_master  # upp_source is "master"


# --- optional files missing -----------------------------------------------


def test_missing_run_json_excluded_and_master_are_handled_leniently(tmp_path):
    data_dir = _make_data_dir(tmp_path, 2025)
    # No run.json, no excluded.csv, no item_master.csv.

    path = build_report(data_dir, tmp_path / "out", 2025, 25)
    wb = load_workbook(path, data_only=False)

    assert wb.sheetnames == ["Top 25", "All items", "Sources"]  # Excluded skipped

    ws = wb["Sources"]
    kv = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    # openpyxl round-trips a cell written with an empty string as None.
    assert kv["item master rows"] is None
    assert kv["top_n"] == 25
    assert "report built at" in kv


# --- geometry at a different top, to catch hardcoded row numbers ----------


def test_geometry_scales_with_top(tmp_path):
    data_dir = _make_data_dir(tmp_path, 2025)
    path = build_report(data_dir, tmp_path / "out", 2025, top=3)

    assert path.name == "Acme Widget Top 3 Items Comparison 2025.xlsx"
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames[0] == "Top 3"
    ws = wb["Top 3"]

    last = 9  # 6 + 3
    assert ws["B4"].value == "Top 3 Items By Quantity"
    assert ws["B9"].value == 3  # last data row, rank 3
    assert ws["K9"].value == "=F9*E9"

    label_row, totals_row, priced_row, basket_row = 10, 11, 12, 13
    assert ws[f"K{label_row}"].value == "Office Depot Total Cost"
    assert ws[f"K{totals_row}"].value == f"=SUM(K7:K{last})"
    assert ws[f"F{priced_row}"].value == f'=COUNT(F7:F{last})&" / 3"'

    basket_formula = ws[f"K{basket_row}"].value
    # All three top items are fully priced by all four vendors in the fixture.
    assert basket_formula == "=SUM(K7,K8,K9)"

    legend_row = basket_row + 2
    assert ws[f"B{legend_row}"].value == "Lowest price for row item"

    cf = ws.conditional_formatting
    ranges = {rng.sqref for rng in cf}
    assert ranges == {f"F7:I{last}", f"K7:N{last}"}
