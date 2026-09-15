"""Builds the leadership workbook: the Top N comparison plus supporting sheets.

Mirrors the 2025 deliverable's Top N page (``docs/spec.md`` section 4.5) with
three deliberate differences: total-cost cells are formulas instead of typed
numbers, the row-minimum highlight is a real conditional-formatting rule
instead of hand fill, and a Comparable-basket row totals only the items every
vendor could price. ``All items``, ``Excluded`` and ``Sources`` make the
reasoning one click away.
"""
from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

# The firm's workbooks (including the 2025 FINAL in inbox/) carry
# ``<font family="34">``, past what openpyxl's default validation allows.
# Widen it once, at import time of this module, so anything that reads a
# workbook afterwards - a later ``load_workbook`` call, including in tests -
# tolerates it too.
openpyxl.styles.fonts.Font.family.max = 99

from .errors import SupplytrackError
from .ingest import (
    CARRY_COLUMNS,
    CARRY_LABELS,
    CARRY_WIDTH_DEFAULT,
    CARRY_WIDTHS,
    INTEGER_COLUMNS,
    carry_rows,
    classify_exports,
)
from .prices import VENDORS, load_prices
from .xlsx import is_canonical_int

PRICE_FORMAT = "$#,##0.00"
QTY_FORMAT = "#,##0"

STATUS_TEXT = {
    "not_available": "not available",
    "discontinued": "Discontinued",
    "unpriced": "not priced",
}

# Light, distinct fills per vendor, reused for the header cell in both the
# unit-price band (F..I) and the total-cost band (K..N).
VENDOR_HEADER_FILLS = {
    "Office Depot": PatternFill(fill_type="solid", fgColor="DDEBF7"),
    "Preferred": PatternFill(fill_type="solid", fgColor="E2EFDA"),
    "Amazon": PatternFill(fill_type="solid", fgColor="FCE4D6"),
    "Staples": PatternFill(fill_type="solid", fgColor="E4DFEC"),
}

LEGEND_COLOR = "FFF2CC"
# Ordinary cell fill (legend cell, column-J divider): solid fill paints with
# the foreground color.
LEGEND_CELL_FILL = PatternFill(fill_type="solid", fgColor=LEGEND_COLOR)
# Conditional-formatting (dxf) fill: Excel's differential-style fill paints
# a solid fill from the *background* color, not the foreground - using
# fgColor here would add a rule that never visibly highlights anything.
CF_HIGHLIGHT_FILL = PatternFill(fill_type="solid", bgColor=LEGEND_COLOR)

TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)

PRICE_COLS = ["F", "G", "H", "I"]
TOTAL_COLS = ["K", "L", "M", "N"]

# Where each carry sheet's rows come from, relative to the data directory, and
# the order the three sheets are written in. One table, so writing the workbook
# and reading it back with `import-report` cannot disagree about either.
CARRY_SOURCES = {
    "item_master": ("item_master.csv", False),
    "prices": ("prices.csv", True),
    "prices_retired": ("prices_retired.csv", True),
}

# Printed on the Sources sheet rather than above a carry sheet's header row:
# the classifier looks for the header in row 1, so nothing may sit above it.
CARRY_NOTE = (
    "The Item master, Prices and Prices retired sheets are read back next year. "
    "Do not edit them by hand."
)


def build_report(data_dir: Path, out_dir: Path, year: int, top: int) -> Path:
    """Write ``out_dir / f"Acme Widget Top {top} Items Comparison {year}.xlsx"`` and return its path."""
    data_dir = Path(data_dir)
    out_dir = Path(out_dir)
    year_dir = data_dir / str(year)

    ranked_rows = _read_csv_rows(year_dir / "ranked.csv")
    if ranked_rows is None:
        raise SupplytrackError(
            f"{year_dir / 'ranked.csv'} not found. Run `supplytrack rank --year {year}` first."
        )
    top_rows = sorted(ranked_rows, key=lambda r: int(r["rank"]))[:top]

    excluded_rows = _read_csv_rows(year_dir / "excluded.csv")
    run_info = _read_json(year_dir / "run.json")
    carry = {
        kind: _read_csv_rows(_carry_source(data_dir, year, kind)) or []
        for kind in CARRY_SOURCES
    }
    master_row_count = _count_master_rows(data_dir / "item_master.csv")

    prices = load_prices(data_dir, year, top)

    wb = Workbook()
    wb.remove(wb.active)

    _build_top_sheet(wb, top_rows, prices, year, top)
    _build_table_sheet(wb, "All items", ranked_rows)
    if excluded_rows is not None:
        _build_table_sheet(wb, "Excluded", excluded_rows)
    unpriced_count = sum(1 for cell in prices.values() if cell.status == "unpriced")
    _build_sources_sheet(wb, run_info, master_row_count, ranked_rows, top, unpriced_count)
    for kind in CARRY_SOURCES:
        _build_carry_sheet(wb, kind, carry[kind])

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"Acme Widget Top {top} Items Comparison {year}.xlsx"
    wb.save(out_path)
    return out_path


def _read_csv_rows(path: Path) -> list[dict[str, str]] | None:
    if not path.exists():
        return None
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _read_json(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def _count_master_rows(path: Path) -> int | None:
    if not path.exists():
        return None
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return sum(1 for _ in csv.DictReader(fh))


def _build_top_sheet(wb, top_rows: list[dict[str, str]], prices, year: int, top: int):
    ws = wb.create_sheet(f"Top {top}")
    first = 7
    last = 6 + top

    ws.merge_cells("B4:N4")
    ws["B4"] = f"Top {top} Items By Quantity"
    ws["B4"].font = TITLE_FONT

    ws.merge_cells("F5:I5")
    ws["F5"] = "Unit Price (per each)"
    ws.merge_cells("K5:N5")
    ws["K5"] = f"Total Cost at {year} volume"

    ws["B6"] = "Ranking"
    ws["C6"] = "Description"
    ws["D6"] = "Pack/Size"
    ws["E6"] = "Quantity"
    for col in ("B", "C", "D", "E"):
        ws[f"{col}6"].font = BOLD
    for col, vendor in zip(PRICE_COLS, VENDORS):
        cell = ws[f"{col}6"]
        cell.value = vendor
        cell.font = BOLD
        cell.fill = VENDOR_HEADER_FILLS[vendor]
    for col, vendor in zip(TOTAL_COLS, VENDORS):
        cell = ws[f"{col}6"]
        cell.value = vendor
        cell.font = BOLD
        cell.fill = VENDOR_HEADER_FILLS[vendor]

    for offset, row in enumerate(top_rows):
        r = first + offset
        canonical_name = row["canonical_name"]
        unit_label = (row.get("unit_label") or "").strip()

        ws[f"B{r}"] = int(row["rank"])
        c_cell = ws[f"C{r}"]
        c_cell.value = canonical_name
        c_cell.alignment = Alignment(wrap_text=True)
        ws[f"D{r}"] = f"1 {unit_label}" if unit_label else "1"
        e_cell = ws[f"E{r}"]
        e_cell.value = int(row["eaches"])
        e_cell.number_format = QTY_FORMAT

        for price_col, total_col, vendor in zip(PRICE_COLS, TOTAL_COLS, VENDORS):
            price_cell = ws[f"{price_col}{r}"]
            total_cell = ws[f"{total_col}{r}"]
            cell_data = prices.get((canonical_name, vendor))
            if cell_data is not None and cell_data.status == "priced" and cell_data.unit_price is not None:
                price_cell.value = float(cell_data.unit_price)
                price_cell.number_format = PRICE_FORMAT
                total_cell.value = f"={price_col}{r}*E{r}"
                total_cell.number_format = PRICE_FORMAT
            else:
                status = cell_data.status if cell_data is not None else "not_available"
                text = STATUS_TEXT.get(status, status)
                price_cell.value = text
                total_cell.value = text

    # Column J is a visual divider between the unit-price and total-cost
    # bands, for the header row and every data row.
    for r in range(6, last + 1):
        ws[f"J{r}"].fill = LEGEND_CELL_FILL

    label_row = last + 1
    totals_row = last + 2
    priced_row = last + 3
    basket_row = last + 4

    for col, vendor in zip(TOTAL_COLS, VENDORS):
        cell = ws[f"{col}{label_row}"]
        cell.value = f"{vendor} Total Cost"
        cell.font = BOLD

    ws[f"C{totals_row}"] = "Total Cost"
    ws[f"C{totals_row}"].font = BOLD
    for col in TOTAL_COLS:
        cell = ws[f"{col}{totals_row}"]
        cell.value = f"=SUM({col}{first}:{col}{last})"
        cell.number_format = PRICE_FORMAT

    ws[f"C{priced_row}"] = "Priced items"
    ws[f"C{priced_row}"].font = BOLD
    for col in PRICE_COLS:
        ws[f"{col}{priced_row}"] = f'=COUNT({col}{first}:{col}{last})&" / {top}"'

    qualifying_rows = [
        first + offset
        for offset, row in enumerate(top_rows)
        if all(
            (cell_data := prices.get((row["canonical_name"], vendor))) is not None
            and cell_data.status == "priced"
            and cell_data.unit_price is not None
            for vendor in VENDORS
        )
    ]

    ws[f"C{basket_row}"] = "Comparable basket"
    ws[f"C{basket_row}"].font = BOLD
    for col in TOTAL_COLS:
        cell = ws[f"{col}{basket_row}"]
        if qualifying_rows:
            cell.value = "=SUM(" + ",".join(f"{col}{r}" for r in qualifying_rows) + ")"
        else:
            cell.value = 0
        cell.number_format = PRICE_FORMAT

    legend_row = basket_row + 2
    ws[f"B{legend_row}"] = "Lowest price for row item"
    ws[f"B{legend_row}"].fill = LEGEND_CELL_FILL

    ws.conditional_formatting.add(
        f"F{first}:I{last}",
        FormulaRule(
            formula=[f"AND(ISNUMBER(F{first}),F{first}=MIN($F{first}:$I{first}))"],
            fill=CF_HIGHLIGHT_FILL,
        ),
    )
    ws.conditional_formatting.add(
        f"K{first}:N{last}",
        FormulaRule(
            formula=[f"AND(ISNUMBER(K{first}),K{first}=MIN($K{first}:$N{first}))"],
            fill=CF_HIGHLIGHT_FILL,
        ),
    )

    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 11
    for col in PRICE_COLS:
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["J"].width = 2
    for col in TOTAL_COLS:
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["C"].alignment = Alignment(wrap_text=True)

    ws.freeze_panes = "C7"
    return ws


def _carry_source(data_dir: Path, year: int, kind: str) -> Path:
    name, per_year = CARRY_SOURCES[kind]
    return (Path(data_dir) / str(year) / name) if per_year else (Path(data_dir) / name)


def _write_carry_cell(cell, text: str, column: str) -> None:
    """One carry-sheet cell, typed so the CSV survives the round trip.

    Everything is written as text except the whole-number columns
    (:data:`ingest.INTEGER_COLUMNS`), so a unit price typed as ``1.250`` comes
    back spelled that way instead of as ``1.25``, and a date stays the ISO
    string the rest of the pipeline writes rather than becoming a date cell
    with a display format of its own. An integer column is only written as a
    number when its text is already that integer's one spelling - ``007`` stays
    text, because ``7`` would not be the same file.

    The ``data_type`` line is not decoration: openpyxl reads a string starting
    with ``=`` as a formula, and a formula cell read back with ``data_only``
    is ``None``, so a note somebody wrote as ``="0000"`` would be lost. The
    JavaScript port has no such rule, so without this the two would disagree.
    """
    if text == "":
        return
    if column in INTEGER_COLUMNS and is_canonical_int(text):
        cell.value = int(text)
        return
    cell.value = text
    cell.data_type = "s"


def _build_carry_sheet(wb, kind: str, rows: list[dict[str, str]]):
    """One of the three sheets that carry a data file forward into next year.

    The header row is row 1 with nothing above it, because that is where the
    sheet classifier looks for it; the guidance that these sheets are machine
    read lives on the Sources sheet instead (:data:`CARRY_NOTE`). The sheet is
    written even when there is nothing in it, so next year's upload finds the
    same three sheets whether or not any prices were retired this year.
    """
    columns = CARRY_COLUMNS[kind]
    ws = wb.create_sheet(CARRY_LABELS[kind])
    for col_idx, header in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = BOLD
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(columns, start=1):
            _write_carry_cell(
                ws.cell(row=row_idx, column=col_idx), str(row.get(header, "") or ""), header
            )
    for col_idx, header in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = CARRY_WIDTHS.get(header, CARRY_WIDTH_DEFAULT)
    ws.freeze_panes = "A2"
    return ws


def import_report(data_dir: Path, year: int, path: Path, force: bool = False):
    """Pull the item master and the price sheets back out of a report workbook.

    This is the other half of the round trip: the workbook the office manager
    keeps is the only file she has to find next year, and this turns it back
    into the three CSVs the pipeline works from. Every target is checked before
    anything is written, so a run that would overwrite one file does not leave
    the other two already replaced.

    Returns ``(written, missing)``: ``(path, row count, label)`` per file
    written, and the labels of the carry sheets the workbook did not have.
    """
    data_dir = Path(data_dir)
    path = Path(path)
    recognised, _ = classify_exports([path], require_export=False)
    found = {}
    for sheet in recognised:
        if sheet.kind in CARRY_SOURCES and sheet.kind not in found:
            found[sheet.kind] = sheet
    if not found:
        raise SupplytrackError(
            f"{path.name} carries none of the sheets this reads: "
            + ", ".join(CARRY_LABELS.values())
            + ". Point it at a report workbook this tool built."
        )

    targets = {kind: _carry_source(data_dir, year, kind) for kind in found}
    if not force:
        clashes = [t for t in targets.values() if t.exists() and t.read_text(encoding="utf-8-sig").strip()]
        if clashes:
            raise SupplytrackError(
                f"{len(clashes)} file(s) already exist and would be overwritten: "
                + ", ".join(str(t) for t in clashes)
                + ". Move them aside, or pass --force to replace them."
            )

    written = []
    for kind, sheet in found.items():
        rows = carry_rows(sheet)
        target = targets[kind]
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=CARRY_COLUMNS[kind])
            writer.writeheader()
            for row in rows:
                writer.writerow({c: row.get(c, "") for c in CARRY_COLUMNS[kind]})
        written.append((target, len(rows), CARRY_LABELS[kind]))
    missing = [CARRY_LABELS[k] for k in CARRY_SOURCES if k not in found]
    return written, missing


def carry_forward(data_dir: Path, year: int, paths: list[Path]):
    """Fill in whichever data files are missing from a report workbook among ``paths``.

    Only ever writes a file the data directory does not already have, so a
    workbook handed to ``run`` alongside the exports never overwrites work in
    progress. Returns ``(path, row count, label, where)`` per file filled in.
    """
    paths = [Path(p) for p in paths]
    if not paths:
        return []
    wanted = {
        kind: target
        for kind, target in ((k, _carry_source(data_dir, year, k)) for k in CARRY_SOURCES)
        if not (target.exists() and target.read_text(encoding="utf-8-sig").strip())
    }
    if not wanted:
        return []
    try:
        recognised, _ = classify_exports(paths, require_export=False)
    except SupplytrackError:
        # Nothing here to carry forward. Whatever is wrong with these files is
        # ingest's to report, in the words it already uses.
        return []

    filled = []
    seen: set[str] = set()
    for sheet in recognised:
        if sheet.kind not in wanted or sheet.kind in seen:
            continue
        seen.add(sheet.kind)
        rows = carry_rows(sheet)
        target = wanted[sheet.kind]
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=CARRY_COLUMNS[sheet.kind])
            writer.writeheader()
            for row in rows:
                writer.writerow({c: row.get(c, "") for c in CARRY_COLUMNS[sheet.kind]})
        filled.append((target, len(rows), CARRY_LABELS[sheet.kind], sheet.label))
    return filled


def _build_table_sheet(wb, name: str, rows: list[dict[str, str]]):
    ws = wb.create_sheet(name)
    headers = list(rows[0].keys()) if rows else []
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = BOLD
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(header, ""))
    if headers:
        last_col = get_column_letter(len(headers))
        last_row = len(rows) + 1
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.freeze_panes = "A2"
    return ws


def _build_sources_sheet(
    wb,
    run_info: dict,
    master_row_count: int | None,
    ranked_rows: list[dict[str, str]],
    top: int,
    unpriced_count: int = 0,
):
    ws = wb.create_sheet("Sources")
    r = 1

    def write_pair(key, value):
        nonlocal r
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=value)
        r += 1

    for key, value in run_info.items():
        if isinstance(value, list):
            value = "; ".join(str(v) for v in value)
        write_pair(key, value)

    write_pair("item master rows", master_row_count if master_row_count is not None else "")

    non_master = [
        row["canonical_name"] for row in ranked_rows if (row.get("upp_source") or "") != "master"
    ]
    write_pair("items with upp_source != master", "; ".join(non_master))

    write_pair("top_n", top)
    write_pair("unpriced cells", unpriced_count)
    write_pair("report built at", datetime.now(timezone.utc).isoformat())
    # Last, so nothing is inserted above the run.json rows a reader walks.
    write_pair("keep this workbook", CARRY_NOTE)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 70
    return ws
