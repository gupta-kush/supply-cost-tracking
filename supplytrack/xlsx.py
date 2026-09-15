"""Reading vendor exports: a tolerant workbook loader, header matching, cell coercion.

Both vendor exports are hand-exported spreadsheets, so they are untidy in ways
that have nothing to do with their content: headers carry stray trailing
spaces, quantities come back as ``2.0``, dates arrive sometimes as real dates
and sometimes as text, prices sometimes carry a dollar sign, and the Amazon
file declares ``<font family="34">``, which openpyxl rejects by default.

Everything that copes with that lives here, so the pipeline modules work with
clean strings, ints and decimals and never see a spreadsheet quirk.
"""
from __future__ import annotations

import csv
import datetime as dt
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl

from .errors import SupplytrackError

_FAMILY_PATCHED = False


def _patch_font_family() -> None:
    """Widen the allowed range of the font ``family`` attribute.

    openpyxl validates ``family`` against the OOXML numbering, which stops well
    below the value the 2025 Amazon export carries (``<font family="34">``).
    That makes an otherwise readable file raise on load. The value is
    presentational and this package never writes these workbooks back, so
    raising the ceiling loses nothing and gains a file that opens.
    """
    global _FAMILY_PATCHED
    if _FAMILY_PATCHED:
        return
    try:
        from openpyxl.styles.fonts import Font

        descriptor = Font.__dict__.get("family")
        if descriptor is not None and hasattr(descriptor, "max"):
            descriptor.max = 1024
    except Exception:  # pragma: no cover - a newer openpyxl may drop the bound
        pass
    _FAMILY_PATCHED = True


def load_workbook_safe(path: Path, data_only: bool = True) -> openpyxl.Workbook:
    """Open a workbook, tolerating the font-family value the exports carry."""
    _patch_font_family()
    try:
        return openpyxl.load_workbook(path, data_only=data_only, read_only=False)
    except SupplytrackError:
        raise
    except Exception as exc:
        raise SupplytrackError(
            f"Could not open the spreadsheet {Path(path).name}: {exc}. "
            "If it came straight off a vendor site, open it in Excel once, save it, and retry."
        ) from exc


def norm_header(h: str) -> str:
    """Strip, case-fold and collapse internal whitespace, for header matching.

    The exports carry trailing spaces on some headers and Excel occasionally
    turns a wrapped header into one with a double space, so matching on the raw
    string fails for a reason nobody can see by looking at the file.
    """
    return re.sub(r"\s+", " ", str(h or "").strip()).casefold()


# How far down a sheet the header row is looked for. A sheet somebody pasted
# an export into often carries a title line or a blank row above the headers,
# so row 1 cannot be assumed; ten rows is far enough to cover that and short
# enough that a sheet of data never looks like a header by accident.
MAX_HEADER_SCAN = 10


def read_table(path: Path, sheet: str | None = None) -> tuple[list[str], list[list]]:
    """Read a .xlsx or .csv table as (normalised headers, rows).

    The first sheet is used unless ``sheet`` names another. The first row is
    the header row and comes back through :func:`norm_header`, so callers match
    on normalised names only. Blank rows are dropped and short rows are padded
    to the header width.
    """
    path = Path(path)
    rows = _raw_grid(path, sheet)
    if not rows:
        raise SupplytrackError(f"{path.name} is empty - there is no header row to read.")
    return normalise_grid(rows, 0)


def normalise_grid(rows: list[list], header_row: int = 0) -> tuple[list[str], list[list]]:
    """Split a raw grid into (normalised headers, body) at ``header_row``.

    Rows above the header row are dropped, blank rows are dropped and short
    rows are padded to the header width. Shared by :func:`read_table` and the
    sheet scanner so a file read either way comes out the same.
    """
    headers = [norm_header(c) for c in rows[header_row]]
    width = len(headers)
    body: list[list] = []
    for row in rows[header_row + 1 :]:
        if not any(str(c).strip() for c in row if c is not None):
            continue
        row = list(row)
        if len(row) < width:
            row += [None] * (width - len(row))
        body.append(row)
    return headers, body


def sheet_grids(path: Path) -> list[tuple[str, list[list]]]:
    """Every sheet of a workbook, or the one table of a .csv, as raw grids.

    Returns ``(sheet name, rows)`` pairs with nothing normalised: the caller
    decides which row is the header. A .csv is a single table and comes back
    with an empty sheet name, because it has no sheet to name.
    """
    path = Path(path)
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return [("", _read_csv(path))]
    if suffix not in (".xlsx", ".xlsm"):
        raise SupplytrackError(
            f"{path.name} has an extension this tool does not read. "
            "Export the order history as .xlsx or .csv."
        )
    if not path.exists():
        raise SupplytrackError(f"File not found: {path}")
    wb = load_workbook_safe(path, data_only=True)
    try:
        return [
            (name, [list(row) for row in wb[name].iter_rows(values_only=True)])
            for name in wb.sheetnames
        ]
    finally:
        wb.close()


def scan_headers(grid: list[list], limit: int = MAX_HEADER_SCAN) -> list[list[str]]:
    """The first ``limit`` rows of a grid, each normalised as a header row."""
    return [[norm_header(c) for c in row] for row in grid[:limit]]


def _raw_grid(path: Path, sheet: str | None) -> list[list]:
    if not path.exists():
        raise SupplytrackError(f"File not found: {path}")
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    raise SupplytrackError(
        f"{path.name} has an extension this tool does not read. "
        "Export the order history as .xlsx or .csv."
    )


def _read_csv(path: Path) -> list[list]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return [list(r) for r in csv.reader(fh)]


def _read_xlsx(path: Path, sheet: str | None) -> list[list]:
    wb = load_workbook_safe(path, data_only=True)
    try:
        if sheet is None:
            ws = wb[wb.sheetnames[0]]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise SupplytrackError(
                f"{path.name} has no sheet named {sheet!r}. It has: {', '.join(wb.sheetnames)}."
            )
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def header_index(
    headers: list[str],
    where: str,
    required: list[str],
    optional: list[str] | None = None,
) -> dict[str, int]:
    """Map the columns this tool reads to their indexes, or say which are missing.

    ``required`` and ``optional`` are given already normalised. Required
    columns must all be there; optional ones are mapped when present and left
    out when not, so an export that dropped a column the tool only copies
    through still reads. One error listing every missing column beats failing
    on the first one: the usual cause is the wrong export having been
    downloaded, and the full list makes that obvious at a glance.
    """
    index = {h: i for i, h in enumerate(headers) if h}
    missing = [h for h in required if h not in index]
    if missing:
        raise SupplytrackError(
            f"{where} is missing {len(missing)} required column(s): "
            + ", ".join(missing)
            + ". Check that this is the full order-history export and not a filtered view."
        )
    mapped = {h: index[h] for h in required}
    for h in optional or []:
        if h in index:
            mapped[h] = index[h]
    return mapped


def cell(row: list, index: dict[str, int], header: str) -> Any:
    """Read one cell by normalised header name, tolerating a short row."""
    i = index.get(header)
    if i is None or i >= len(row):
        return None
    return row[i]


def clean_text(value: Any) -> str:
    """A cell as trimmed text, with Excel's formula wrapper removed.

    Amazon writes some identifier columns as the formula string ``="0000"`` so
    Excel keeps the leading zeros. Anything shaped like that is unwrapped.
    """
    if value is None:
        return ""
    text = str(value).strip()
    m = re.fullmatch(r'=\s*"(.*)"', text, flags=re.DOTALL)
    if m:
        return m.group(1).strip()
    return text


_DATE_FORMATS = (
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%Y/%m/%d",
    "%b %d, %Y",
    "%B %d, %Y",
    "%d-%b-%Y",
)


def coerce_date(value: Any, *, where: str = "") -> str:
    """Return an ISO ``YYYY-MM-DD`` date from a date cell or a text date.

    Whether a date arrives as a real date or as text depends on how the export
    was downloaded, so both are accepted. An unreadable value raises rather
    than quietly becoming blank, because a wrong date moves a line into or out
    of the reporting year.
    """
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raw = clean_text(value)
    if not raw:
        raise SupplytrackError(f"Missing order date{_at(where)}.")
    text = raw.split("T")[0].strip()
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    try:  # e.g. "2025-01-05 00:00:00"
        return dt.datetime.fromisoformat(raw).date().isoformat()
    except ValueError:
        pass
    raise SupplytrackError(
        f"Could not read {raw!r} as an order date{_at(where)}. "
        "Expected a date cell or a date such as 2025-01-05 or 1/5/2025."
    )


def coerce_int(value: Any, *, where: str = "") -> int:
    """Return a whole number from a quantity cell.

    openpyxl hands back ``2.0`` for a quantity typed as 2, and some exports
    quote it as text. A blank or unreadable quantity raises naming the row: a
    quantity that silently became zero is exactly the quiet loss the
    line-count check exists to catch.
    """
    if isinstance(value, bool):
        raise SupplytrackError(f"Expected a quantity but found a true/false value{_at(where)}.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise SupplytrackError(f"Quantity {value} is not a whole number{_at(where)}.")
    text = clean_text(value).replace(",", "")
    if not text:
        raise SupplytrackError(f"Missing quantity{_at(where)}.")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise SupplytrackError(f"Could not read {text!r} as a quantity{_at(where)}.") from exc
    if number != number.to_integral_value():
        raise SupplytrackError(f"Quantity {text} is not a whole number{_at(where)}.")
    return int(number)


def coerce_decimal(value: Any, *, where: str = "") -> str:
    """Return a money amount as a plain decimal string, or "" when blank.

    The Preferred export carries no prices at all, so blank is normal here and
    must not be an error. Dollar signs, thousands commas and parenthesised
    negatives are stripped, because hand-exported sheets carry all three.
    """
    if isinstance(value, bool):
        raise SupplytrackError(f"Expected an amount but found a true/false value{_at(where)}.")
    if isinstance(value, (int, float, Decimal)):
        return _trim(Decimal(str(value)))
    text = clean_text(value)
    if not text:
        return ""
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace(",", "").strip()
    if not text:
        return ""
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise SupplytrackError(f"Could not read {text!r} as an amount{_at(where)}.") from exc
    return _trim(-number if negative else number)


def _trim(number: Decimal) -> str:
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _at(where: str) -> str:
    return f" ({where})" if where else ""
